"""Unit tests for attachments/processors/ -- one class per processor, each
against a real, generated file of that type (not a hand-rolled byte
fixture), so a corruption/format-detail assumption gets caught the same way
a real upload would."""

from __future__ import annotations

import io
import json

from PIL import Image

from attachments.models import Attachment
from attachments.processors.csv_processor import CsvProcessor
from attachments.processors.docx_processor import DocxProcessor
from attachments.processors.image_processor import ImageProcessor
from attachments.processors.json_processor import JsonProcessor
from attachments.processors.pdf_processor import PdfProcessor
from attachments.processors.pptx_processor import PptxProcessor
from attachments.processors.registry import get_processor_for
from attachments.processors.text_processor import TextProcessor
from attachments.processors.xlsx_processor import XlsxProcessor
from config.settings import Settings


def _settings(**overrides) -> Settings:
    return Settings(_env_file=None, **overrides)


def _attachment(tmp_path, filename: str, media_type: str, content: bytes) -> Attachment:
    path = tmp_path / filename
    path.write_bytes(content)
    return Attachment(
        original_filename=filename,
        safe_filename=filename,
        media_type=media_type,
        extension=path.suffix,
        size_bytes=len(content),
        sha256="test-hash",
        local_path=str(path),
    )


class TestRegistry:
    def test_returns_none_for_an_unregistered_media_type(self):
        att = Attachment(
            original_filename="x.bin",
            safe_filename="x.bin",
            media_type="application/octet-stream",
            extension=".bin",
            size_bytes=1,
            sha256="x",
        )
        assert get_processor_for(att) is None


class TestTextProcessor:
    def test_extracts_plain_text(self, tmp_path):
        att = _attachment(tmp_path, "notes.txt", "text/plain", b"line one\nline two")
        outcome = TextProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert "line one" in outcome.extracted_text
        assert outcome.metadata["line_count"] == 2

    def test_truncates_past_the_char_limit(self, tmp_path):
        att = _attachment(tmp_path, "big.txt", "text/plain", ("x" * 1000).encode())
        outcome = TextProcessor().process(att, _settings(max_attachment_text_chars=100))
        assert len(outcome.extracted_text) == 100
        assert outcome.metadata["truncated"] is True

    def test_falls_back_to_latin1_for_invalid_utf8(self, tmp_path):
        att = _attachment(tmp_path, "bad_encoding.txt", "text/plain", b"\xff\xfe not utf8")
        outcome = TextProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert "fallback" in outcome.metadata["encoding"]


class TestJsonProcessor:
    def test_pretty_prints_valid_json(self, tmp_path):
        content = json.dumps({"b": 1, "a": 2}).encode()
        att = _attachment(tmp_path, "data.json", "application/json", content)
        outcome = JsonProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert outcome.metadata["top_level_type"] == "dict"

    def test_rejects_invalid_json_with_a_useful_error(self, tmp_path):
        att = _attachment(tmp_path, "bad.json", "application/json", b"{not valid json")
        outcome = JsonProcessor().process(att, _settings())
        assert outcome.status == "failed"
        assert "Invalid JSON" in outcome.error


class TestCsvProcessor:
    def test_reads_headers_and_rows(self, tmp_path):
        content = b"name,age\nAlice,30\nBob,25\n"
        att = _attachment(tmp_path, "people.csv", "text/csv", content)
        outcome = CsvProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert outcome.metadata["column_count"] == 2
        assert outcome.metadata["row_count"] == 2
        assert outcome.extracted_tables[0]["headers"] == ["name", "age"]

    def test_truncates_and_notes_row_limit(self, tmp_path):
        rows = "\n".join(f"row{i},{i}" for i in range(20))
        content = f"name,value\n{rows}\n".encode()
        att = _attachment(tmp_path, "many.csv", "text/csv", content)
        outcome = CsvProcessor().process(att, _settings(max_attachment_spreadsheet_rows=5))
        assert outcome.status == "succeeded"
        assert outcome.metadata["truncated"] is True
        assert outcome.metadata["rows_included"] == 5
        assert "Truncated" in outcome.extracted_text

    def test_empty_csv_fails_cleanly(self, tmp_path):
        att = _attachment(tmp_path, "empty.csv", "text/csv", b"")
        outcome = CsvProcessor().process(att, _settings())
        assert outcome.status == "failed"


class TestPdfProcessor:
    def _make_pdf_bytes(self, text: str) -> bytes:
        from reportlab.pdfgen import canvas

        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer)
        pdf.drawString(100, 750, text)
        pdf.save()
        return buffer.getvalue()

    def test_extracts_text_with_page_marker(self, tmp_path):
        content = self._make_pdf_bytes("Total revenue was 42000 dollars.")
        att = _attachment(tmp_path, "report.pdf", "application/pdf", content)
        outcome = PdfProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert "42000" in outcome.extracted_text
        assert "[Page 1]" in outcome.extracted_text
        assert outcome.metadata["page_count"] == 1

    def test_corrupt_pdf_fails_cleanly(self, tmp_path):
        att = _attachment(tmp_path, "bad.pdf", "application/pdf", b"%PDF-1.4\nnot really a pdf")
        outcome = PdfProcessor().process(att, _settings())
        assert outcome.status == "failed"

    def test_enforces_page_limit(self, tmp_path):
        from reportlab.pdfgen import canvas

        buffer = io.BytesIO()
        pdf = canvas.Canvas(buffer)
        for i in range(5):
            pdf.drawString(100, 750, f"Page {i}")
            pdf.showPage()
        pdf.save()
        att = _attachment(tmp_path, "multi.pdf", "application/pdf", buffer.getvalue())
        outcome = PdfProcessor().process(att, _settings(max_attachment_document_pages=2))
        assert outcome.status == "failed"
        assert "page" in outcome.error.lower()

    def test_rejects_a_pdf_with_embedded_javascript(self, tmp_path):
        """Wiring test for attachments.pdf_safety.check_pdf_safety -- a
        real PDF built with pypdf's own `add_js` writer helper, genuinely
        exercising the /Names/JavaScript catalog structure a malicious PDF
        would use, never a checked-in malware sample."""
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_js('app.alert("hello");')
        buffer = io.BytesIO()
        writer.write(buffer)

        att = _attachment(tmp_path, "malicious.pdf", "application/pdf", buffer.getvalue())
        outcome = PdfProcessor().process(att, _settings())
        assert outcome.status == "failed"
        assert "javascript" in outcome.error.lower()

    def test_rejects_a_pdf_with_an_embedded_file(self, tmp_path):
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        writer.add_attachment("note.txt", b"hello world")
        buffer = io.BytesIO()
        writer.write(buffer)

        att = _attachment(tmp_path, "carrier.pdf", "application/pdf", buffer.getvalue())
        outcome = PdfProcessor().process(att, _settings())
        assert outcome.status == "failed"
        assert "embedded files" in outcome.error.lower()


class TestDocxProcessor:
    def test_extracts_paragraphs_and_tables(self, tmp_path):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Executive summary text.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Score"
        table.cell(1, 0).text = "Alice"
        table.cell(1, 1).text = "95"
        buffer = io.BytesIO()
        doc.save(buffer)

        att = _attachment(
            tmp_path,
            "report.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            buffer.getvalue(),
        )
        outcome = DocxProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert "Executive summary" in outcome.extracted_text
        assert outcome.metadata["table_count"] == 1
        assert outcome.extracted_tables[0]["rows"][1] == ["Alice", "95"]

    def test_corrupt_docx_fails_cleanly(self, tmp_path):
        att = _attachment(
            tmp_path,
            "bad.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            b"not a zip file",
        )
        outcome = DocxProcessor().process(att, _settings())
        assert outcome.status == "failed"

    def test_zip_decompression_bomb_guard_is_wired_in(self, tmp_path):
        """Wiring test for attachments.zip_safety.check_zip_safety -- a
        real, ordinary, valid DOCX (a legitimate small ZIP archive), with
        the configured uncompressed-size limit set artificially low so the
        guard actually trips, the same "real file + tight limit" pattern
        test_enforces_page_limit above already uses for PDF pages. Proves
        the processor calls the guard and honors a rejection -- not a claim
        that this file itself is malicious."""
        from docx import Document

        doc = Document()
        doc.add_paragraph("Executive summary text.")
        buffer = io.BytesIO()
        doc.save(buffer)

        att = _attachment(
            tmp_path,
            "report.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            buffer.getvalue(),
        )
        outcome = DocxProcessor().process(att, _settings(max_attachment_zip_uncompressed_bytes=10))
        assert outcome.status == "failed"
        assert "decompression bomb" in outcome.error.lower()


class TestXlsxProcessor:
    def test_extracts_sheet_data(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Region", "Total"])
        ws.append(["West", 1000])
        ws.append(["East", 2000])
        buffer = io.BytesIO()
        wb.save(buffer)

        att = _attachment(
            tmp_path,
            "sales.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            buffer.getvalue(),
        )
        outcome = XlsxProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert outcome.metadata["sheet_names"] == ["Sales"]
        assert "West" in outcome.extracted_text
        assert outcome.extracted_tables[0]["headers"] == ["Region", "Total"]

    def test_corrupt_xlsx_fails_cleanly(self, tmp_path):
        att = _attachment(
            tmp_path,
            "bad.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            b"not a zip file",
        )
        outcome = XlsxProcessor().process(att, _settings())
        assert outcome.status == "failed"

    def test_zip_decompression_bomb_guard_is_wired_in(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Region", "Total"])
        buffer = io.BytesIO()
        wb.save(buffer)

        att = _attachment(
            tmp_path,
            "sales.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            buffer.getvalue(),
        )
        outcome = XlsxProcessor().process(att, _settings(max_attachment_zip_uncompressed_bytes=10))
        assert outcome.status == "failed"
        assert "decompression bomb" in outcome.error.lower()


class TestPptxProcessor:
    def test_extracts_slide_title_and_body(self, tmp_path):
        from pptx import Presentation

        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        slide.shapes.title.text = "Q3 Results"
        body = slide.placeholders[1]
        body.text_frame.text = "Revenue grew 12%."
        buffer = io.BytesIO()
        pres.save(buffer)

        att = _attachment(
            tmp_path,
            "deck.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            buffer.getvalue(),
        )
        outcome = PptxProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert "[Slide 1]" in outcome.extracted_text
        assert "Q3 Results" in outcome.extracted_text
        assert "Revenue grew 12%" in outcome.extracted_text
        assert outcome.metadata["slide_count"] == 1

    def test_corrupt_pptx_fails_cleanly(self, tmp_path):
        att = _attachment(
            tmp_path,
            "bad.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            b"not a zip file",
        )
        outcome = PptxProcessor().process(att, _settings())
        assert outcome.status == "failed"

    def test_zip_decompression_bomb_guard_is_wired_in(self, tmp_path):
        from pptx import Presentation

        pres = Presentation()
        slide = pres.slides.add_slide(pres.slide_layouts[1])
        slide.shapes.title.text = "Q3 Results"
        buffer = io.BytesIO()
        pres.save(buffer)

        att = _attachment(
            tmp_path,
            "deck.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            buffer.getvalue(),
        )
        outcome = PptxProcessor().process(att, _settings(max_attachment_zip_uncompressed_bytes=10))
        assert outcome.status == "failed"
        assert "decompression bomb" in outcome.error.lower()


class TestImageProcessor:
    def test_produces_a_data_url(self, tmp_path):
        buffer = io.BytesIO()
        Image.new("RGB", (50, 50), color="green").save(buffer, format="PNG")
        att = _attachment(tmp_path, "square.png", "image/png", buffer.getvalue())
        outcome = ImageProcessor().process(att, _settings())
        assert outcome.status == "succeeded"
        assert outcome.image_data_url.startswith("data:image/png;base64,")

    def test_downscales_oversized_images(self, tmp_path):
        buffer = io.BytesIO()
        Image.new("RGB", (4000, 100), color="red").save(buffer, format="PNG")
        att = _attachment(tmp_path, "wide.png", "image/png", buffer.getvalue())
        outcome = ImageProcessor().process(att, _settings(max_attachment_image_dimension_px=500))
        assert outcome.status == "succeeded"

    def test_corrupt_image_fails_cleanly(self, tmp_path):
        att = _attachment(tmp_path, "bad.png", "image/png", b"\x89PNG\r\n\x1a\n" + b"garbage")
        outcome = ImageProcessor().process(att, _settings())
        assert outcome.status == "failed"
