"""Processor for Excel (.xlsx) attachments -- reads every sheet (bounded per
sheet by `Settings.max_attachment_spreadsheet_rows`) via `openpyxl`, in
read-only mode (`read_only=True`) so a large workbook isn't fully loaded
into memory just to read its first N rows. Deliberately not pandas -- same
reasoning `csv_processor.py`'s docstring gives (this codebase's documented
pandas/Python-3.14 datetime segfault, CLAUDE.md's "Python 3.14 gotchas")."""

from __future__ import annotations

import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from attachments.models import Attachment
from attachments.processors.base import ProcessingOutcome
from attachments.zip_safety import ZipSafetyError, check_zip_safety
from config.settings import Settings


def _cell_to_str(value: object) -> str:
    return "" if value is None else str(value)


def _render_sheet(sheet_name: str, rows: list[list[str]], truncated: bool, total_rows: int) -> str:
    lines = [f"[Sheet: {sheet_name}]"]
    if rows:
        header, *body = rows
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join("---" for _ in header) + " |")
        for row in body:
            padded = row + [""] * (len(header) - len(row))
            lines.append("| " + " | ".join(padded[: len(header)]) + " |")
    else:
        lines.append("(empty sheet)")
    if truncated:
        lines.append(f"[Truncated: showing {len(rows) - 1} of {total_rows} data row(s).]")
    return "\n".join(lines)


class XlsxProcessor:
    source_type = "xlsx"

    def can_process(self, attachment: Attachment) -> bool:
        return (
            attachment.media_type
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    def process(self, attachment: Attachment, settings: Settings) -> ProcessingOutcome:
        path = Path(attachment.local_path)  # type: ignore[arg-type]

        try:
            check_zip_safety(
                path.read_bytes(),
                max_uncompressed_bytes=settings.max_attachment_zip_uncompressed_bytes,
                max_entries=settings.max_attachment_zip_entries,
            )
        except ZipSafetyError as exc:
            return ProcessingOutcome(status="failed", error=str(exc))

        try:
            workbook = load_workbook(str(path), data_only=True, read_only=True)
        except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as exc:
            return ProcessingOutcome(
                status="failed",
                error=f"This XLSX file appears to be corrupted or invalid: {exc}.",
            )

        max_rows = settings.max_attachment_spreadsheet_rows
        sheet_sections: list[str] = []
        extracted_tables: list[dict] = []
        truncated_any = False

        try:
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows: list[list[str]] = []
                total_rows = 0
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index == 0:
                        rows.append([_cell_to_str(v) for v in row])
                        continue
                    total_rows += 1
                    if len(rows) <= max_rows:
                        rows.append([_cell_to_str(v) for v in row])

                truncated = total_rows > max_rows
                truncated_any = truncated_any or truncated
                sheet_sections.append(_render_sheet(sheet_name, rows, truncated, total_rows))
                if rows:
                    extracted_tables.append(
                        {"sheet_name": sheet_name, "headers": rows[0], "rows": rows[1:]}
                    )
        finally:
            workbook.close()

        combined_text = "\n\n".join(sheet_sections)
        text_truncated = len(combined_text) > settings.max_attachment_text_chars
        if text_truncated:
            combined_text = combined_text[: settings.max_attachment_text_chars]

        return ProcessingOutcome(
            status="succeeded",
            extracted_text=combined_text,
            extracted_tables=extracted_tables or None,
            metadata={
                "sheet_names": workbook.sheetnames,
                "row_limit_reached": truncated_any,
                "truncated": text_truncated,
            },
        )
