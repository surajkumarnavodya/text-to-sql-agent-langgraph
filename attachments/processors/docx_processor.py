"""Processor for Word (.docx) attachments -- extracts paragraphs and tables
via `python-docx`, preserving basic document structure (paragraph order,
table rows/columns) rather than flattening everything into one blob."""

from __future__ import annotations

import zipfile
from pathlib import Path

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from attachments.models import Attachment
from attachments.processors.base import ProcessingOutcome
from attachments.zip_safety import ZipSafetyError, check_zip_safety
from config.settings import Settings


class DocxProcessor:
    source_type = "docx"

    def can_process(self, attachment: Attachment) -> bool:
        return (
            attachment.media_type
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    def process(self, attachment: Attachment, settings: Settings) -> ProcessingOutcome:
        path = Path(attachment.local_path)  # type: ignore[arg-type]

        try:
            check_zip_safety(
                path.read_bytes(),
                max_uncompressed_bytes=settings.max_attachment_zip_uncompressed_bytes,
                max_entries=settings.max_attachment_zip_entries,
            )
        except ZipSafetyError as exc:
            return ProcessingOutcome(status="failed", error=str(exc))

        try:
            document = Document(str(path))
        except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
            return ProcessingOutcome(
                status="failed",
                error=f"This DOCX file appears to be corrupted or invalid: {exc}.",
            )

        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

        extracted_tables: list[dict] = []
        table_sections: list[str] = []
        for table_index, table in enumerate(document.tables):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            extracted_tables.append({"table_index": table_index, "rows": rows})
            rendered_rows = "\n".join(" | ".join(row) for row in rows)
            table_sections.append(f"[Table {table_index + 1}]\n{rendered_rows}")

        parts = ["\n\n".join(paragraphs)]
        if table_sections:
            parts.append("\n\n".join(table_sections))
        combined_text = "\n\n".join(part for part in parts if part.strip())

        truncated = len(combined_text) > settings.max_attachment_text_chars
        if truncated:
            combined_text = combined_text[: settings.max_attachment_text_chars]

        if not combined_text.strip():
            return ProcessingOutcome(
                status="failed",
                error="No extractable text or tables were found in this document.",
            )

        return ProcessingOutcome(
            status="succeeded",
            extracted_text=combined_text,
            extracted_tables=extracted_tables or None,
            metadata={
                "paragraph_count": len(paragraphs),
                "table_count": len(document.tables),
                "truncated": truncated,
            },
        )
