"""Processor for PowerPoint (.pptx) attachments -- extracts each slide's
title and body text, with slide numbers preserved, via `python-pptx`."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.exc import PackageNotFoundError

from attachments.models import Attachment
from attachments.processors.base import ProcessingOutcome
from attachments.zip_safety import ZipSafetyError, check_zip_safety
from config.settings import Settings


def _slide_title(slide) -> str | None:  # noqa: ANN001 - pptx.slide.Slide, no public type export
    if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
        text = slide.shapes.title.text_frame.text.strip()
        return text or None
    return None


def _slide_text(slide) -> list[str]:  # noqa: ANN001 - see _slide_title
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            parts.append(shape.text_frame.text.strip())
    return parts


class PptxProcessor:
    source_type = "pptx"

    def can_process(self, attachment: Attachment) -> bool:
        return (
            attachment.media_type
            == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        )

    def process(self, attachment: Attachment, settings: Settings) -> ProcessingOutcome:
        path = Path(attachment.local_path)  # type: ignore[arg-type]

        try:
            check_zip_safety(
                path.read_bytes(),
                max_uncompressed_bytes=settings.max_attachment_zip_uncompressed_bytes,
                max_entries=settings.max_attachment_zip_entries,
            )
        except ZipSafetyError as exc:
            return ProcessingOutcome(status="failed", error=str(exc))

        try:
            presentation = Presentation(str(path))
        except (PackageNotFoundError, KeyError, ValueError) as exc:
            return ProcessingOutcome(
                status="failed",
                error=f"This PPTX file appears to be corrupted or invalid: {exc}.",
            )

        slide_sections: list[str] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            title = _slide_title(slide)
            text_parts = _slide_text(slide)
            # The title shape's own text is already included in text_parts
            # (it's a normal text-frame shape too) -- deduplicate so it
            # doesn't render twice, once as the heading and once in the body.
            body_parts = [t for t in text_parts if t != title]
            header = f"[Slide {slide_number}]" + (f" {title}" if title else "")
            body = "\n".join(body_parts)
            slide_sections.append(f"{header}\n{body}".strip())

        combined_text = "\n\n".join(section for section in slide_sections if section)
        truncated = len(combined_text) > settings.max_attachment_text_chars
        if truncated:
            combined_text = combined_text[: settings.max_attachment_text_chars]

        if not combined_text.strip():
            return ProcessingOutcome(
                status="failed",
                error="No extractable text was found in this presentation.",
            )

        return ProcessingOutcome(
            status="succeeded",
            extracted_text=combined_text,
            metadata={
                "slide_count": len(presentation.slides),
                "truncated": truncated,
            },
        )
